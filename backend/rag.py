import os
import re
import faiss
import numpy as np
import fitz
from PIL import Image
from pptx import Presentation
from io import BytesIO
from google import genai
from google.genai import types

# AQ.* API keys require v1beta endpoint for gemini-2.5-flash
GENAI_HTTP_OPTIONS = {"api_version": "v1beta"}


# ──────────────────────────────────────────────
# Deterministic skill extraction — reads the resume's own
# "Skills" / "Tech Stack" section instead of asking an LLM to infer skills.
# ──────────────────────────────────────────────
SKILL_SECTION_HEADERS = {
    "technical skills", "skills and tools", "skills & tools", "tech stack",
    "technologies", "core competencies", "key skills", "skill set", "skills",
    "tools & technologies", "tools and technologies", "technical skillset",
    "technical expertise", "areas of expertise",
}

SECTION_STOP_HEADERS = {
    "experience", "work experience", "professional experience", "employment",
    "employment history", "education", "projects", "project experience",
    "certifications", "certificates", "achievements", "awards", "publications",
    "extracurricular", "extracurriculars", "activities", "summary", "objective",
    "profile", "contact", "references", "internship", "internships",
    "leadership", "volunteer", "volunteering", "languages", "interests",
    "hobbies", "about", "about me", "personal projects",
}

_SPLIT_RE = re.compile(r'[,|;•·•]+')
_HEADING_SPLIT_RE = re.compile(r'\s*&\s*|\s+and\s+')


def _normalize_heading(line: str) -> str:
    return re.sub(r'[^a-z& ]', '', line.strip().lower()).strip()


def _heading_parts(line: str) -> list[str]:
    """Return normalized heading fragments if this line looks like a standalone
    section heading (short, no ':' content) — else []. Splits combined
    headings like 'Achievements & Certifications' into individual parts so
    each can be checked against the keyword sets."""
    if ':' in line:
        return []
    norm = _normalize_heading(line)
    if not norm or len(norm) > 40:
        return []
    words = norm.split()
    if len(words) > 6:
        return []
    parts = [p.strip() for p in _HEADING_SPLIT_RE.split(norm) if p.strip()]
    return parts or [norm]


def _matches_heading(line: str, keywords: set) -> bool:
    return any(part in keywords for part in _heading_parts(line))


def extract_skills_from_resume(resume_text: str) -> list[str]:
    """Parse the literal Skills / Tech Stack section of a resume.
    Returns [] if no such heading is found."""
    lines = resume_text.splitlines()
    n = len(lines)

    start_idx = None
    for i, line in enumerate(lines):
        if _matches_heading(line, SKILL_SECTION_HEADERS):
            start_idx = i + 1
            break
    if start_idx is None:
        return []

    end_idx = n
    for j in range(start_idx, n):
        if _matches_heading(lines[j], SECTION_STOP_HEADERS):
            end_idx = j
            break

    skills: list[str] = []
    seen = set()
    for line in lines[start_idx:end_idx]:
        line = line.strip()
        if not line:
            continue
        if ':' in line:
            line = line.split(':', 1)[1]
        for tok in _SPLIT_RE.split(line):
            tok = re.sub(r'\s+', ' ', tok.strip(' -\t')).strip()
            if not tok or len(tok) > 50:
                continue
            key = tok.lower()
            if key in seen:
                continue
            seen.add(key)
            skills.append(tok)
    return skills


# ──────────────────────────────────────────────
# Text chunker with overlap (used for chatbot RAG)
# ──────────────────────────────────────────────
def chunk_text(text: str, chunk_size: int = 1000, overlap: int = 200) -> list[str]:
    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunks.append(text[start:end])
        start += (chunk_size - overlap)
    return [c.strip() for c in chunks if c.strip()]


# ──────────────────────────────────────────────
# RAG Engine
# ──────────────────────────────────────────────
class RAGEngine:
    EMBEDDING_MODEL = "gemini-embedding-001"
    GENERATION_MODEL = "gemini-2.5-flash"

    def __init__(self, api_key: str):
        self.client = genai.Client(api_key=api_key, http_options=GENAI_HTTP_OPTIONS)
        self.index = None
        self.chunks: list[str] = []
        self.raw_text = ""
        self.extracted_images = []

    # ── Step 1: Extract raw text ──────────────
    def extract_text(self, file_bytes: bytes, filename: str, content_type: str) -> str:
        text = ""
        filename_lower = filename.lower()
        try:
            if "pdf" in content_type or filename_lower.endswith(".pdf"):
                doc = fitz.open(stream=file_bytes, filetype="pdf")
                for page in doc:
                    text += page.get_text() + "\n"
                    # Extract images
                    for img in page.get_images(full=True):
                        xref = img[0]
                        base_image = doc.extract_image(xref)
                        image_bytes = base_image["image"]
                        try:
                            pil_img = Image.open(BytesIO(image_bytes))
                            if pil_img.mode not in ('RGB', 'L'):
                                pil_img = pil_img.convert('RGB')
                            self.extracted_images.append(pil_img)
                        except Exception:
                            pass

            elif "presentation" in content_type or filename_lower.endswith(".pptx"):
                prs = Presentation(BytesIO(file_bytes))
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if hasattr(shape, "text") and shape.text.strip():
                            text += shape.text.strip() + "\n"
            else:
                text = file_bytes.decode("utf-8", errors="ignore")

        except Exception as e:
            print(f"[RAG] Text extraction error for '{filename}': {e}")
            raise ValueError(f"Could not parse '{filename}'. File may be corrupted or unsupported.")

        return text.strip()

    # ── Step 2: Direct Summarization (NO RAG) ─
    def generate_summary(self, text: str) -> str:
        """Summarize the entire extracted text directly using Gemini."""
        self.raw_text = text
        
        # Trim if it exceeds Gemini's context window (highly unlikely for student notes)
        context = text[:600000]

        print(f"[RAG] Generating summary directly from entire text ({len(context)} chars). No RAG used.")

        prompt = f"""
You are an expert AI learning assistant. A student has uploaded their study notes.
Carefully read the entire text below and produce a structured study guide.

Notes:
---
{context}
---

Respond ONLY with a valid JSON object in exactly this format (no markdown, no extra text):
{{
  "summary": "A cohesive executive summary combining all the key ideas from the notes (150-250 words).",
  "concepts": [
    {{
      "title": "Concept name",
      "description": "Clear, concise explanation of the concept."
    }}
  ],
  "revisionNotes": [
    "Short, actionable bullet point for rapid revision.",
    "Another bullet point."
  ]
}}
"""
        payload = [prompt]
        if hasattr(self, "extracted_images") and self.extracted_images:
            payload.extend(self.extracted_images[:15])

        try:
            result = self.client.models.generate_content(
                model=self.GENERATION_MODEL,
                contents=payload,
            )
            text_out = result.text.strip()
        except Exception as e:
            raise ValueError(f"Gemini generation failed: {e}")

        # Strip markdown fences if present
        if text_out.startswith("```json"):
            text_out = text_out[7:]
        elif text_out.startswith("```"):
            text_out = text_out[3:]
        if text_out.endswith("```"):
            text_out = text_out[:-3]

        return text_out.strip()

    # ── Step 3: Setup Vector Index for Chatbot ──
    def setup_chat_index(self, text: str):
        """Chunk and embed the full text to prepare the FAISS index for chatbot Q&A."""
        self.chunks = chunk_text(text)
        if not self.chunks:
            print("[Chat RAG] No text chunks produced. Chatbot index setup skipped.")
            return

        print(f"[Chat RAG] Preparing vector index: embedding {len(self.chunks)} chunks...")
        embeddings = []

        for i, chunk in enumerate(self.chunks):
            try:
                response = self.client.models.embed_content(
                    model=self.EMBEDDING_MODEL,
                    contents=chunk,
                    config=types.EmbedContentConfig(task_type="RETRIEVAL_DOCUMENT"),
                )
                embeddings.append(response.embeddings[0].values)
            except Exception as e:
                print(f"[Chat RAG] Failed to embed chunk {i}: {e}")

        if not embeddings:
            print("[Chat RAG] Warning: Chatbot embedding failed for all chunks.")
            return

        embeddings_np = np.array(embeddings, dtype="float32")
        dimension = embeddings_np.shape[1]

        # Build FAISS flat L2 index
        self.index = faiss.IndexFlatL2(dimension)
        self.index.add(embeddings_np)
        print(f"[Chat RAG] FAISS index built with {self.index.ntotal} vectors (dim={dimension})")

    # ── Step 4: Chatbot Q&A (RAG-based) ────────
    def answer_chat_question(self, query: str, history: list = None) -> str:
        """Answer user questions about their notes by retrieving relevant chunks first, keeping track of history."""
        if not self.chunks:
            return "No document has been uploaded yet. Please upload a document first."

        # If we have a vector index, retrieve relevant chunks
        if self.index is not None:
            try:
                q_response = self.client.models.embed_content(
                    model=self.EMBEDDING_MODEL,
                    contents=query,
                    config=types.EmbedContentConfig(task_type="RETRIEVAL_QUERY"),
                )
                query_embedding = np.array([q_response.embeddings[0].values], dtype="float32")
                
                # Retrieve top 5 most relevant chunks
                k = min(5, len(self.chunks))
                distances, indices = self.index.search(query_embedding, k)
                
                retrieved_chunks = [self.chunks[int(idx)] for idx in indices[0] if int(idx) < len(self.chunks)]
                context = "\n\n---\n\n".join(retrieved_chunks)
                print(f"[Chat RAG] Retrieved {len(retrieved_chunks)} relevant chunks for query: '{query}'")
            except Exception as e:
                print(f"[Chat RAG] Retrieval failed: {e}. Falling back to full text context.")
                context = self.raw_text[:100000]
        else:
            context = self.raw_text[:100000]

        # Format conversation history
        history_str = ""
        if history:
            for msg in history:
                role = "Student" if msg.get("role") == "user" else "Tutor"
                history_str += f"{role}: {msg.get('text')}\n"

        # Build the chat prompt
        prompt = f"""
You are an expert AI tutor. A student is asking questions about their uploaded study notes.
Answer their question accurately, clearly, and concisely, using ONLY the context provided below.
Maintain a helpful and educational tone.

If the answer cannot be found in the context, politely let the student know you don't have that information.

Context from student's notes:
---
{context}
---

Conversation History:
{history_str}
Student: {query}
Tutor:
"""
        try:
            result = self.client.models.generate_content(
                model=self.GENERATION_MODEL,
                contents=prompt,
            )
            return result.text.strip()
        except Exception as e:
            return f"Error generating answer: {str(e)}"

    # ── Step 5: Quiz Generation ────────────────
    def generate_quiz(self, difficulty: str, question_count: int = 10) -> str:
        """Generate a quiz (Medium/Advanced) based on the extracted notes."""
        if not self.raw_text:
            raise ValueError("No text available to generate a quiz. Please upload notes first.")

        context = self.raw_text[:600000]

        # Calculate split: ~80% MCQ, ~20% subjective
        num_subjective = max(2, round(question_count * 0.2))
        num_mcq = question_count - num_subjective

        prompt = f"""
You are an expert AI tutor. Based on the student's study notes provided below, generate a {difficulty.upper()} difficulty quiz.
The quiz should test their understanding of the core concepts in the notes.
You MUST include EXACTLY {question_count} questions: {num_mcq} Multiple Choice Questions (MCQs) and {num_subjective} Subjective (or Coding, if applicable) questions.
If images are provided, use the visual information in them to formulate relevant questions.

Notes:
---
{context}
---

Respond ONLY with a valid JSON object in exactly this format (no markdown, no extra text):
{{
  "difficulty": "{difficulty}",
  "topic": "Broad topic name of the notes (e.g. Cryptography, Operating Systems, Machine Learning)",
  "questions": [
    {{
      "type": "mcq",
      "skill": "Name of the specific skill or concept tested",
      "question": "The question text",
      "options": ["Option A", "Option B", "Option C", "Option D"],
      "answer": "The correct option text exactly as it appears in the options array",
      "explanation": "Brief explanation of why this is correct."
    }},
    {{
      "type": "subjective",
      "skill": "Name of the specific skill or concept tested",
      "question": "The subjective or coding question text",
      "answer": "A model answer or key points expected in the answer.",
      "explanation": "Explanation or grading rubric."
    }}
  ]
}}
"""
        payload = [prompt]
        if hasattr(self, "extracted_images") and self.extracted_images:
            payload.extend(self.extracted_images[:15])

        try:
            result = self.client.models.generate_content(
                model=self.GENERATION_MODEL,
                contents=payload,
            )
            text_out = result.text.strip()
        except Exception as e:
            raise ValueError(f"Gemini generation failed: {e}")

        # Strip markdown fences if present
        if text_out.startswith("```json"):
            text_out = text_out[7:]
        elif text_out.startswith("```"):
            text_out = text_out[3:]
        if text_out.endswith("```"):
            text_out = text_out[:-3]

        return text_out.strip()

    # ── Step 6: Career Analysis (Resume → Gaps + Projects) ──
    def analyze_career(self, resume_text: str, desired_role: str, current_skills: list[str]) -> str:
        """Analyze a resume against a desired role, given a skill list already
        extracted deterministically from the resume's Skills/Tech Stack section.
        Returns JSON with skillsToLearn, readinessScore, and categorized projects."""

        context = resume_text[:400000]
        skills_str = ", ".join(current_skills) if current_skills else "(none listed)"

        prompt = f"""
You are an expert career coach and technical hiring manager. A candidate has uploaded their resume and wants to become a "{desired_role}".

The candidate's current skills, extracted directly from the Skills/Tech Stack section of their resume, are:
{skills_str}

Using the full resume below only for extra context, your task:
1. Identify the skills that a strong "{desired_role}" must have that are NOT in the candidate's current skills list above.
2. Score the candidate's overall readiness for the "{desired_role}" role on a 0-100 scale, based on how many of the critical, role-defining skills are already present (from the current skills list above) versus missing. Be realistic and strict — a candidate missing core fundamentals for the role (e.g. no data structures/algorithms for an SDE role) should score low even if they have many unrelated skills.
3. Suggest 9 hands-on projects (3 beginner, 3 intermediate, 3 advanced) that will help build the missing skills and significantly strengthen the resume for this role.

Resume:
---
{context}
---

Respond ONLY with a valid JSON object in exactly this format (no markdown, no extra text):
{{
  "desiredRole": "{desired_role}",
  "skillsToLearn": [
    "Missing skill needed for the role"
  ],
  "readinessScore": 0,
  "scoreExplanation": "1-2 sentence honest explanation of the score, naming specific strong areas and specific critical gaps.",
  "projects": [
    {{
      "level": "beginner",
      "title": "Project title",
      "description": "What the project does and why it's relevant to the role.",
      "skillsBuilt": ["Skill 1", "Skill 2"],
      "howToBuild": "A concise, practical suggestion on how to approach building this project — key steps, tools, or resources to use.",
      "githubRepo": "https://github.com/owner/repo — a real, well-known GitHub repository that the user can study as a reference implementation or starting point for this project. Only include repos you are highly confident exist."
    }},
    {{
      "level": "intermediate",
      "title": "Project title",
      "description": "What the project does and why it's relevant to the role.",
      "skillsBuilt": ["Skill 1", "Skill 2"],
      "howToBuild": "A concise, practical suggestion on how to approach building this project — key steps, tools, or resources to use.",
      "githubRepo": "https://github.com/owner/repo — a real, well-known GitHub repository that the user can study as a reference implementation or starting point for this project. Only include repos you are highly confident exist."
    }},
    {{
      "level": "advanced",
      "title": "Project title",
      "description": "What the project does and why it's relevant to the role.",
      "skillsBuilt": ["Skill 1", "Skill 2"],
      "howToBuild": "A concise, practical suggestion on how to approach building this project — key steps, tools, or resources to use.",
      "githubRepo": "https://github.com/owner/repo — a real, well-known GitHub repository that the user can study as a reference implementation or starting point for this project. Only include repos you are highly confident exist."
    }}
  ]
}}

Include exactly 3 beginner, 3 intermediate, and 3 advanced projects in the projects array.
For githubRepo, provide ONLY the URL (starting with https://github.com/) with no extra text or description after it.
"readinessScore" must be a plain integer between 0 and 100 (no quotes, no % sign).
"""
        try:
            result = self.client.models.generate_content(
                model=self.GENERATION_MODEL,
                contents=prompt,
            )
            text_out = result.text.strip()
        except Exception as e:
            raise ValueError(f"Gemini career analysis failed: {e}")

        # Strip markdown fences if present
        if text_out.startswith("```json"):
            text_out = text_out[7:]
        elif text_out.startswith("```"):
            text_out = text_out[3:]
        if text_out.endswith("```"):
            text_out = text_out[:-3]

        return text_out.strip()
